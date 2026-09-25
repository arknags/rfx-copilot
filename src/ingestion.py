import hashlib
import shutil
from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import load_workbook
from PIL import Image


SUPPORTED_FILE_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}

MAX_FILE_SIZE_BYTES = 45 * 1024 * 1024  # 45 MB


class IngestionError(RuntimeError):
    """Raised when a source file cannot be ingested safely."""


@dataclass(frozen=True)
class InboxFile:
    path: Path
    filename: str
    extension: str
    mime_type: str
    file_size_bytes: int
    sha256_checksum: str


@dataclass(frozen=True)
class PreparedSourceDocument:
    """
    A validated source file copied into data/processed/<rfx_id>/.
    """

    original_path: Path
    processed_path: Path
    filename: str
    extension: str
    mime_type: str
    file_size_bytes: int
    sha256_checksum: str
    baseline_content: str


def calculate_sha256(file_path: Path) -> str:
    """Calculate a SHA-256 checksum without loading the whole file into memory."""

    digest = hashlib.sha256()

    with open(file_path, "rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def inspect_inbox_file(file_path: Path) -> InboxFile:
    """Validate and describe one inbox file."""

    if not file_path.exists() or not file_path.is_file():
        raise IngestionError(f"File does not exist: {file_path}")

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_FILE_TYPES:
        raise IngestionError(
            f"Unsupported file type '{extension}' for file {file_path.name}."
        )

    file_size_bytes = file_path.stat().st_size

    if file_size_bytes == 0:
        raise IngestionError(f"File is empty: {file_path.name}")

    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        raise IngestionError(
            f"{file_path.name} is larger than the {MAX_FILE_SIZE_BYTES // 1024 // 1024} MB limit."
        )

    return InboxFile(
        path=file_path,
        filename=file_path.name,
        extension=extension,
        mime_type=SUPPORTED_FILE_TYPES[extension],
        file_size_bytes=file_size_bytes,
        sha256_checksum=calculate_sha256(file_path),
    )


def list_inbox_files(inbox_dir: Path) -> list[InboxFile]:
    """Return all supported files in data/inbox."""

    if not inbox_dir.exists():
        inbox_dir.mkdir(parents=True, exist_ok=True)

    inbox_files = []

    for file_path in sorted(inbox_dir.iterdir()):
        if file_path.is_file() and not file_path.name.startswith("."):
            try:
                inbox_files.append(inspect_inbox_file(file_path))
            except IngestionError:
                # Unsupported files are ignored by the inbox list.
                # The UI can separately report them later if needed.
                continue

    return inbox_files


def _safe_rfx_directory_name(rfx_id: str) -> str:
    """Prevent path separators or special characters in an RFx folder name."""

    return "".join(
        character
        for character in rfx_id
        if character.isalnum() or character in {"-", "_"}
    )


def copy_to_processed(
    *,
    inbox_file: InboxFile,
    processed_root: Path,
    rfx_id: str,
) -> Path:
    """
    Copy the original source file to data/processed/<rfx_id>/.

    The inbox file is retained unchanged.
    """

    safe_rfx_id = _safe_rfx_directory_name(rfx_id)

    if not safe_rfx_id:
        raise IngestionError("Invalid RFx ID for processed-file storage.")

    destination_dir = processed_root / safe_rfx_id
    destination_dir.mkdir(parents=True, exist_ok=True)

    destination_path = (
        destination_dir
        / f"{inbox_file.sha256_checksum[:12]}_{inbox_file.filename}"
    )

    if not destination_path.exists():
        shutil.copy2(inbox_file.path, destination_path)

    return destination_path


def extract_excel_content(file_path: Path) -> str:
    """Extract worksheet cells with sheet and cell references."""

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    lines = []

    for worksheet in workbook.worksheets:
        lines.append(f"=== Worksheet: {worksheet.title} ===")

        for row in worksheet.iter_rows():
            values = []

            for cell in row:
                if cell.value is not None:
                    values.append(f"{cell.coordinate}={cell.value}")

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


def extract_pdf_content(file_path: Path) -> str:
    """Extract text page by page and retain page references."""

    lines = []

    with pdfplumber.open(file_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            lines.append(f"=== Page {page_number} ===")
            lines.append(text if text.strip() else "[No extractable text found]")

    return "\n".join(lines)


def extract_docx_content(file_path: Path) -> str:
    """Extract paragraphs and tables from a Word document."""

    document = Document(file_path)
    lines = []

    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()

        if text:
            lines.append(f"Paragraph {paragraph_number}: {text}")

    for table_number, table in enumerate(document.tables, start=1):
        lines.append(f"=== Table {table_number} ===")

        for row_number, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(f"Row {row_number}: " + " | ".join(cells))

    return "\n".join(lines)


def extract_text_content(file_path: Path) -> str:
    """Read a plain-text vendor email or quote."""

    return file_path.read_text(encoding="utf-8", errors="replace")


def extract_image_metadata(file_path: Path) -> str:
    """
    Preserve simple image metadata. Gemini receives the actual image separately.
    """

    with Image.open(file_path) as image:
        return (
            f"Image file: {file_path.name}\n"
            f"Format: {image.format}\n"
            f"Dimensions: {image.width} x {image.height}\n"
            f"Mode: {image.mode}\n"
            "Note: Use the original image as the authoritative visual source."
        )


def extract_baseline_content(inbox_file: InboxFile) -> str:
    """
    Extract deterministic baseline content before Gemini interpretation.
    """

    if inbox_file.extension == ".xlsx":
        return extract_excel_content(inbox_file.path)

    if inbox_file.extension == ".pdf":
        return extract_pdf_content(inbox_file.path)

    if inbox_file.extension == ".docx":
        return extract_docx_content(inbox_file.path)

    if inbox_file.extension == ".txt":
        return extract_text_content(inbox_file.path)

    if inbox_file.extension in {".jpg", ".jpeg"}:
        return extract_image_metadata(inbox_file.path)

    raise IngestionError(
        f"No baseline extractor exists for {inbox_file.extension}."
    )


def prepare_source_document(
    *,
    file_path: Path,
    processed_root: Path,
    rfx_id: str,
) -> PreparedSourceDocument:
    """
    Validate, preserve, and baseline-extract one inbox document.

    This function does not call Gemini and does not write to SQLite.
    """

    inbox_file = inspect_inbox_file(file_path)

    processed_path = copy_to_processed(
        inbox_file=inbox_file,
        processed_root=processed_root,
        rfx_id=rfx_id,
    )

    baseline_content = extract_baseline_content(inbox_file)

    return PreparedSourceDocument(
        original_path=inbox_file.path,
        processed_path=processed_path,
        filename=inbox_file.filename,
        extension=inbox_file.extension,
        mime_type=inbox_file.mime_type,
        file_size_bytes=inbox_file.file_size_bytes,
        sha256_checksum=inbox_file.sha256_checksum,
        baseline_content=baseline_content,
    )